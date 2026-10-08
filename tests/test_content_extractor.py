from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook

import content_extractor


def test_extract_txt_text(tmp_path: Path):
    path = tmp_path / "note.txt"
    path.write_text("PAN card\nPermanent Account Number", encoding="utf-8")
    assert "PAN card" in content_extractor.extract_txt_text(str(path))


def test_extract_docx_text_reads_paragraphs_and_tables(tmp_path: Path):
    path = tmp_path / "marksheet.docx"
    document = Document()
    document.add_paragraph("12th Marksheet")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Anshul"
    document.save(path)

    text = content_extractor.extract_docx_text(str(path))
    assert "12th Marksheet" in text
    assert "Name | Anshul" in text


def test_extract_xlsx_text(tmp_path: Path):
    path = tmp_path / "income.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Income"
    sheet.append(["Name", "Income"])
    sheet.append(["Anshul", 30000])
    workbook.save(path)
    workbook.close()

    text = content_extractor.extract_xlsx_text(str(path))
    assert "[Sheet: Income]" in text
    assert "Anshul | 30000" in text


def test_extract_pdf_text_uses_embedded_text_without_ocr(tmp_path: Path, monkeypatch):
    path = tmp_path / "sample.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "This is a long embedded document text that does not need OCR.")
    document.save(path)
    document.close()

    called = False

    def fail_if_ocr(*args, **kwargs):
        nonlocal called
        called = True
        return ""

    monkeypatch.setattr(content_extractor, "extract_scanned_pdf_text", fail_if_ocr)
    text = content_extractor.extract_pdf_text(str(path))
    assert "embedded document text" in text
    assert called is False


def test_extract_text_wraps_extraction_errors(tmp_path: Path, monkeypatch):
    path = tmp_path / "broken.docx"
    path.write_text("not a real docx", encoding="utf-8")
    monkeypatch.setattr(content_extractor, "extract_docx_text", lambda _: (_ for _ in ()).throw(ValueError("bad file")))
    try:
        content_extractor.extract_text(str(path))
    except RuntimeError as error:
        assert "broken.docx" in str(error)
        assert "bad file" in str(error)
    else:
        raise AssertionError("extract_text should wrap lower-level extraction failures")
