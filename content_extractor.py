from pathlib import Path
from typing import Tuple

import pymupdf
from docx import Document
from openpyxl import load_workbook

from ocr_engine import extract_image_text, extract_scanned_pdf_text


def extract_docx_text(file_path: str) -> str:
    parts = []
    document = Document(file_path)
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for section in document.sections:
        for paragraph in section.header.paragraphs + section.footer.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                parts.append(" | ".join(values))
    return "\n".join(parts).strip()


def extract_xlsx_text(file_path: str) -> str:
    parts = []
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            parts.append(f"[Sheet: {worksheet.title}]")
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    parts.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(parts).strip()


def extract_xls_text(file_path: str) -> str:
    try:
        import xlrd
    except ImportError:
        raise RuntimeError("Legacy .xls support requires xlrd. Install requirements.txt again.")
    workbook = xlrd.open_workbook(file_path, on_demand=True)
    parts = []
    try:
        for sheet in workbook.sheets():
            parts.append(f"[Sheet: {sheet.name}]")
            for row_index in range(sheet.nrows):
                values = [str(value).strip() for value in sheet.row_values(row_index) if str(value).strip()]
                if values:
                    parts.append(" | ".join(values))
    finally:
        workbook.release_resources()
    return "\n".join(parts).strip()


def extract_txt_text(file_path: str) -> str:
    return Path(file_path).read_text(encoding="utf-8", errors="ignore").strip()


def _pdf_text_with_page_fallback(file_path: str) -> Tuple[str, bool]:
    parts = []
    needs_ocr = False
    with pymupdf.open(file_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            if len(text) >= 30:
                parts.append(f"[Page {page_number}]\n{text}")
            else:
                needs_ocr = True
                # Keep any small amount of embedded text; OCR supplements it.
                if text:
                    parts.append(f"[Page {page_number}]\n{text}")
    return "\n\n".join(parts).strip(), needs_ocr


def extract_pdf_text(file_path: str) -> str:
    text, needs_ocr = _pdf_text_with_page_fallback(file_path)
    if not needs_ocr:
        return text
    ocr_text = extract_scanned_pdf_text(file_path, skip_pages_with_text=True)
    if ocr_text:
        return f"{text}\n\n{ocr_text}".strip()
    return text


def extract_text(file_path: str) -> str:
    extension = Path(file_path).suffix.lower()
    try:
        if extension == ".pdf":
            return extract_pdf_text(file_path)
        if extension == ".docx":
            return extract_docx_text(file_path)
        if extension == ".xlsx":
            return extract_xlsx_text(file_path)
        if extension == ".xls":
            return extract_xls_text(file_path)
        if extension == ".txt":
            return extract_txt_text(file_path)
        if extension in {".jpg", ".jpeg", ".png"}:
            return extract_image_text(file_path)
        return ""
    except Exception as error:
        raise RuntimeError(f"Could not extract text from {Path(file_path).name}: {error}") from error
